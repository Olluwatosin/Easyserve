"""Reading a menu out of whatever the venue already has it in.

A new lounge's menu arrives as whatever its designer or consultant produced: a
spreadsheet, a Word document, a PDF laid out for printing. Typing a hundred
lines into a web form is most of a day's work and the kind of job that gets
three-quarters done.

**Nothing here writes anything.** Parsing returns rows for a human to look at,
correct and confirm. That separation is the whole design, and it is what makes
supporting PDF defensible at all:

  * CSV is exact. The columns say what they are.
  * DOCX is usually reliable — a menu in Word is normally a table, and a table
    has cells.
  * PDF is a guess. A PDF is a description of ink on a page, not a list of
    dishes: columns interleave, decorative fonts extract as nonsense, and a
    scanned menu has no text in it whatsoever.

A parser that is right 70% of the time is a large time-saving when a person
reviews it, and a disaster when it writes straight to the menu — forty items at
plausible-looking wrong prices, discovered by a guest. So the confidence of each
row travels with it, and the unreadable ones are reported as unreadable rather
than dropped quietly.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

#: Words that mark a section heading rather than an item. A heading has no
#: price, so most are excluded by that alone; these catch the rest.
_SECTION_HINTS = {
    "starters", "mains", "main course", "sides", "desserts", "grills",
    "cocktails", "mocktails", "spirits", "whisky", "whiskey", "cognac",
    "champagne", "wine", "wines", "beer", "beers", "soft drinks", "shots",
    "small chops", "platters", "menu", "drinks", "food", "extras", "bottles",
}

#: Section headings that imply what the things under them are.
_FOOD_WORDS = {
    "starter", "main", "side", "dessert", "grill", "chop", "platter",
    "kitchen", "food", "pepper soup", "suya", "rice", "swallow", "snack",
}
_DRINK_WORDS = {
    "cocktail", "mocktail", "spirit", "whisky", "whiskey", "cognac", "vodka",
    "gin", "rum", "tequila", "champagne", "wine", "beer", "drink", "shot",
    "bottle", "soft", "juice", "water", "mixer",
}

# ₦45,000 · N45,000 · 45,000 · 45000.00 · 45k
_PRICE = re.compile(
    # The grouped alternative carries its own optional decimals: without them
    # "2,500.50" matched "2,500" and stopped, importing the price fifty kobo
    # light and silently.
    r"(?:₦|N|NGN)?\s*(\d{1,3}(?:[,\s]\d{3})+(?:\.\d{1,2})?|\d+(?:\.\d{1,2})?)\s*(k\b)?",
    re.IGNORECASE,
)


@dataclass
class ParsedItem:
    name: str
    price: float | None = None
    category: str | None = None
    item_type: str = "other"
    unit_cost: float | None = None
    description: str | None = None
    #: "high" for a column that said what it was, "low" for a line we split on
    #: a guess. The review screen sorts the doubtful ones to the top.
    confidence: str = "high"


@dataclass
class ParseResult:
    items: list[ParsedItem] = field(default_factory=list)
    #: Lines that looked like they should be items but could not be read, kept
    #: verbatim so a person can see what was missed rather than wonder.
    unreadable: list[str] = field(default_factory=list)
    source: str = ""

    @property
    def needs_review(self) -> int:
        return sum(1 for i in self.items if i.confidence != "high" or i.price is None)


def parse_price(text: str) -> float | None:
    """Read a Nigerian menu price. Returns None when there isn't one."""
    if not text:
        return None
    m = _PRICE.search(text.replace(" ", " "))
    if not m:
        return None
    raw = m.group(1).replace(",", "").replace(" ", "")
    try:
        value = float(raw)
    except ValueError:
        return None
    if m.group(2):  # "45k"
        value *= 1000
    # A menu price of zero is a parse artefact, not a free dish.
    return value if value > 0 else None


def _looks_like_section(line: str) -> bool:
    stripped = line.strip(" -–—:·.\t")
    if not stripped or len(stripped) > 40:
        return False
    if parse_price(stripped) is not None:
        return False
    low = stripped.lower()
    if low in _SECTION_HINTS or any(h in low for h in _SECTION_HINTS):
        return True
    # ALL CAPS with no price is a heading in almost every menu ever printed.
    return stripped.isupper() and len(stripped.split()) <= 5


def _type_from(category: str | None, name: str) -> str:
    hay = f"{category or ''} {name}".lower()
    if any(w in hay for w in _DRINK_WORDS):
        return "drink"
    if any(w in hay for w in _FOOD_WORDS):
        return "food"
    return "other"


def _split_name_and_price(line: str) -> tuple[str, float | None]:
    """A printed menu line is a name, some leader dots, and a price."""
    price = parse_price(line)
    if price is None:
        return line.strip(" .-–—\t"), None
    # Remove the last price-looking run, which is where the price sits.
    matches = list(_PRICE.finditer(line))
    cut = matches[-1].start() if matches else len(line)
    name = line[:cut].strip(" .-–—:·\t")
    return name, price


# ── CSV ──────────────────────────────────────────────────────────────────────

_HEADER_ALIASES = {
    "name": {"name", "item", "item name", "product", "dish", "description"},
    "price": {"price", "selling price", "amount", "rate", "cost to guest", "price (₦)", "price(n)"},
    "unit_cost": {"cost", "unit cost", "cost price", "buying price", "purchase price"},
    "category": {"category", "section", "group", "type", "menu section"},
    "item_type": {"item type", "kind", "food/drink", "drink/food"},
}


def _match_header(cell: str) -> str | None:
    key = cell.strip().lower().replace("_", " ")
    for field_name, aliases in _HEADER_ALIASES.items():
        if key in aliases:
            return field_name
    return None


def parse_csv(data: bytes) -> ParseResult:
    """Exact, because the columns say what they are."""
    out = ParseResult(source="csv")
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        return out

    header = [_match_header(c) for c in rows[0]]
    has_header = any(h is not None for h in header)
    body = rows[1:] if has_header else rows

    for row in body:
        if not any(c.strip() for c in row):
            continue
        rec: dict[str, str] = {}
        if has_header:
            for idx, field_name in enumerate(header):
                if field_name and idx < len(row):
                    rec[field_name] = row[idx].strip()
        else:
            # No header: assume the two most common columns, in order.
            rec["name"] = row[0].strip()
            if len(row) > 1:
                rec["price"] = row[1].strip()

        name = rec.get("name", "").strip()
        if not name:
            continue
        price = parse_price(rec.get("price", ""))
        category = rec.get("category") or None
        declared = (rec.get("item_type") or "").strip().lower()
        item_type = declared if declared in ("drink", "food", "other") else _type_from(category, name)

        out.items.append(
            ParsedItem(
                name=name,
                price=price,
                category=category,
                item_type=item_type,
                unit_cost=parse_price(rec.get("unit_cost", "")),
                confidence="high" if has_header else "low",
            )
        )
    return out


# ── DOCX ─────────────────────────────────────────────────────────────────────

def parse_docx(data: bytes) -> ParseResult:
    """Tables first — a menu in Word is usually one — then loose paragraphs."""
    out = ParseResult(source="docx")
    try:
        import docx  # python-docx
    except ImportError:  # pragma: no cover - dependency is declared
        raise RuntimeError("Reading Word documents needs python-docx")

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        raise ValueError("That does not look like a Word document")

    current_section: str | None = None

    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            cells = [c for c in cells if c]
            if not cells:
                continue
            if len(cells) == 1:
                if _looks_like_section(cells[0]):
                    current_section = cells[0].strip(" -–—:")
                continue
            name = cells[0]
            price = next((p for p in (parse_price(c) for c in cells[1:]) if p), None)
            if _looks_like_section(name) and price is None:
                current_section = name.strip(" -–—:")
                continue
            out.items.append(
                ParsedItem(
                    name=name,
                    price=price,
                    category=current_section,
                    item_type=_type_from(current_section, name),
                    confidence="high" if price is not None else "low",
                )
            )

    for para in document.paragraphs:
        line = para.text.strip()
        if not line:
            continue
        if _looks_like_section(line):
            current_section = line.strip(" -–—:")
            continue
        name, price = _split_name_and_price(line)
        if not name or price is None:
            continue
        if any(i.name.lower() == name.lower() for i in out.items):
            continue
        out.items.append(
            ParsedItem(
                name=name,
                price=price,
                category=current_section,
                item_type=_type_from(current_section, name),
                confidence="low",
            )
        )
    return out


# ── PDF ──────────────────────────────────────────────────────────────────────

def parse_pdf(data: bytes) -> ParseResult:
    """A guess, and labelled as one.

    A PDF describes ink, not data. Columns interleave, decorative fonts extract
    as nonsense, and a scanned menu contains no text at all — which is reported
    rather than returned as an empty success.
    """
    out = ParseResult(source="pdf")
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - dependency is declared
        raise RuntimeError("Reading PDFs needs pypdf")

    try:
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:
        raise ValueError("That PDF could not be opened — is the file corrupt?")

    if not text.strip():
        raise ValueError(
            "No text could be read from that PDF. If it is a scan or a photo, "
            "the menu will have to be typed or supplied as a spreadsheet."
        )

    current_section: str | None = None
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if _looks_like_section(line):
            current_section = line.strip(" -–—:")
            continue
        name, price = _split_name_and_price(line)
        if not name:
            continue
        if price is None:
            # A line with no price is usually a description or a stray header.
            # Kept only if it is long enough to be worth a person's glance.
            if len(name) > 3 and len(name.split()) <= 8:
                out.unreadable.append(line)
            continue
        if len(name) < 2:
            out.unreadable.append(line)
            continue
        out.items.append(
            ParsedItem(
                name=name,
                price=price,
                category=current_section,
                item_type=_type_from(current_section, name),
                confidence="low",
            )
        )
    return out


# ── Entry point ──────────────────────────────────────────────────────────────

def parse(filename: str, data: bytes) -> ParseResult:
    lower = (filename or "").lower()
    if lower.endswith(".csv") or lower.endswith(".txt"):
        return parse_csv(data)
    if lower.endswith(".docx"):
        return parse_docx(data)
    if lower.endswith(".pdf"):
        return parse_pdf(data)
    if lower.endswith(".doc"):
        raise ValueError(
            "Old .doc files cannot be read. Open it in Word and save as .docx, "
            "or export the menu as a PDF."
        )
    if lower.endswith((".xlsx", ".xls")):
        raise ValueError(
            "Excel files are not read directly. In Excel choose File → Save As "
            "→ CSV, which imports exactly."
        )
    raise ValueError("Upload a CSV, Word (.docx) or PDF file")
